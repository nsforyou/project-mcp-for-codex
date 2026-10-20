#!/usr/bin/env python
# -*- coding: utf-8 -*-
r"""把 mpp_probe.py 导出的 JSON 转成可直接打开的 Excel 任务清单。

字段覆盖：WBS编号、任务名称、父任务、负责人、计划开始、计划完成、
实际开始、实际完成、工期、完成百分比、状态、前置任务、是否里程碑、备注。

用法:
    python scripts/export_excel.py data/template.json [exports/输出.xlsx]
"""
from __future__ import annotations

import json
import sys
from datetime import date, datetime
from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

ROOT = Path(__file__).resolve().parent.parent

HEADERS = [
    ("WBS编号", 12),
    ("层级", 6),
    ("任务名称", 52),
    ("父任务", 30),
    ("负责人", 22),
    ("计划开始", 12),
    ("计划完成", 12),
    ("实际开始", 12),
    ("实际完成", 12),
    ("工期(天)", 9),
    ("完成百分比", 10),
    ("状态", 10),
    ("前置任务", 24),
    ("是否里程碑", 10),
    ("摘要任务", 9),
    ("备注", 40),
]

THIN = Side(style="thin", color="D9D9D9")
BORDER = Border(left=THIN, right=THIN, top=THIN, bottom=THIN)


def day(value: str | None) -> str:
    return value[:10] if value else ""


def derive_status(pct, start: str | None, finish: str | None) -> str:
    today = date.today().isoformat()
    if pct is None:
        pct = 0
    if pct >= 100:
        return "已完成"
    if pct > 0:
        return "进行中"
    if start and day(start) > today:
        return "未开始"
    if finish and day(finish) < today:
        return "已逾期未开始"
    return "未开始"


def pred_text(task: dict) -> str:
    items = []
    for pr in task.get("predecessors", []):
        t = pr.get("type") or ""
        short = {"FINISH_START": "FS", "START_START": "SS", "FINISH_FINISH": "FF", "START_FINISH": "SF"}.get(t, t)
        txt = f"{pr.get('predecessor_id')}({short}"
        lag = (pr.get("lag") or {}).get("text")
        if lag and lag not in ("0.0d", "0d"):
            txt += f"+{lag}"
        items.append(txt + ")")
    return ", ".join(items)


def build(json_path: Path, out_path: Path) -> Path:
    data = json.loads(json_path.read_text(encoding="utf-8"))
    tasks = data["tasks"]
    proj = data["project"]
    stats = data["stats"]

    wb = Workbook()
    ws = wb.active
    ws.title = "任务清单"

    # 标题行
    ws.merge_cells(start_row=1, start_column=1, end_row=1, end_column=len(HEADERS))
    title = ws.cell(row=1, column=1, value=f"{proj['title']}　计划期间 {day(proj['start_date'])} ~ {day(proj['finish_date'])}")
    title.font = Font(bold=True, size=13)
    title.alignment = Alignment(vertical="center")
    ws.row_dimensions[1].height = 24

    header_row = 2
    for col, (name, width) in enumerate(HEADERS, start=1):
        c = ws.cell(row=header_row, column=col, value=name)
        c.font = Font(bold=True, color="FFFFFF")
        c.fill = PatternFill("solid", fgColor="4472C4")
        c.alignment = Alignment(horizontal="center", vertical="center")
        c.border = BORDER
        ws.column_dimensions[get_column_letter(col)].width = width
    ws.freeze_panes = ws.cell(row=header_row + 1, column=3)

    summary_fill = PatternFill("solid", fgColor="DDEBF7")
    milestone_fill = PatternFill("solid", fgColor="FFF2CC")

    r = header_row
    for t in tasks:
        r += 1
        level = t.get("outline_level") or 0
        dur = (t.get("duration") or {}).get("days")
        pct = t.get("percent_complete")
        # Project 保存后会把 0% 存成空值，导出时按 0 处理
        if pct is None:
            pct = 0
        row = [
            t.get("outline_number") or t.get("wbs") or t.get("id"),
            level,
            ("　" * max(level - 1, 0)) + (t.get("name") or ""),
            t.get("parent_name") or "",
            ", ".join(t.get("assignees") or []),
            day(t.get("start")),
            day(t.get("finish")),
            day(t.get("actual_start")),
            day(t.get("actual_finish")),
            dur,
            (pct / 100) if isinstance(pct, (int, float)) else pct,
            derive_status(pct, t.get("start"), t.get("finish")),
            pred_text(t),
            "是" if t.get("is_milestone") else "",
            "是" if t.get("is_summary") else "",
            t.get("notes") or "",
        ]
        for col, value in enumerate(row, start=1):
            c = ws.cell(row=r, column=col, value=value)
            c.border = BORDER
            if col == 11 and isinstance(value, float):
                c.number_format = "0%"
            if col == 12:
                c.alignment = Alignment(horizontal="center")
            if t.get("is_milestone"):
                c.fill = milestone_fill
            elif t.get("is_summary"):
                c.fill = summary_fill
        if t.get("is_summary"):
            for col in range(1, len(HEADERS) + 1):
                ws.cell(row=r, column=col).font = Font(bold=True)

    ws.auto_filter.ref = f"A{header_row}:{get_column_letter(len(HEADERS))}{r}"

    # 第二个工作表：项目信息与统计
    ws2 = wb.create_sheet("项目信息与统计")
    ws2.column_dimensions["A"].width = 24
    ws2.column_dimensions["B"].width = 60
    info = [
        ("源文件", data["source_file"]),
        ("文件大小(KB)", data["file_size_kb"]),
        ("项目标题", proj["title"]),
        ("作者", proj.get("author") or ""),
        ("公司", proj["company"]),
        ("计划开始", day(proj["start_date"])),
        ("计划完成", day(proj["finish_date"])),
        ("默认日历", proj["default_calendar"]),
        ("最后保存", day(proj["last_saved"])),
        ("", ""),
        ("任务总数", stats["task_count"]),
        ("摘要任务", stats["summary_count"]),
        ("明细任务", stats["leaf_count"]),
        ("里程碑", stats["milestone_count"]),
        ("依赖关系(任务侧)", stats["relation_count"]),
        ("依赖关系(文件级)", stats.get("project_relation_count", 0)),
        ("资源数", stats["resource_count"]),
        ("日历数", stats["calendar_count"]),
        ("最大层级", stats["max_outline_level"]),
        ("一级 WBS 数", stats["top_level_count"]),
        ("导出时间", datetime.now().strftime("%Y-%m-%d %H:%M:%S")),
    ]
    for i, (k, v) in enumerate(info, start=1):
        a = ws2.cell(row=i, column=1, value=k)
        a.font = Font(bold=bool(k))
        ws2.cell(row=i, column=2, value=v)

    out_path.parent.mkdir(parents=True, exist_ok=True)
    wb.save(out_path)
    return out_path


def main() -> int:
    if len(sys.argv) < 2:
        print(__doc__)
        return 1
    json_path = Path(sys.argv[1])
    if not json_path.is_absolute():
        json_path = ROOT / json_path
    out_path = Path(sys.argv[2]) if len(sys.argv) > 2 else ROOT / "exports" / (json_path.stem + "_任务清单.xlsx")
    if not out_path.is_absolute():
        out_path = ROOT / out_path
    result = build(json_path, out_path)
    print(f"已生成: {result}")
    return 0


if __name__ == "__main__":
    sys.exit(main())
